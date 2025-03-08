import os
import json
import groq
import dotenv
import streamlit as st
import tempfile
import dateparser
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF

# Load environment variables
dotenv.load_dotenv("touch.env")
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

if not GROQ_API_KEY:
    raise ValueError("GROQ_API_KEY is missing. Please check your touch.env file.")

# Initialize Groq Client
client = groq.Client(api_key=GROQ_API_KEY)

def extract_text_from_file(uploaded_file):
    """Extract text from TXT, PDF, or DOCX files."""
    if uploaded_file.type == "text/plain":
        return uploaded_file.read().decode("utf-8")
    elif uploaded_file.type == "application/pdf":
        pdf_reader = PdfReader(uploaded_file)
        return "\n".join([page.extract_text() for page in pdf_reader.pages if page.extract_text()])
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(uploaded_file)
        return "\n".join([para.text for para in doc.paragraphs])
    return None

def summarize_transcript(transcript):
    """Use Groq API to summarize the meeting transcript."""
    prompt = f"""
    Summarize the following meeting transcript in a concise and informative way.
    Highlight the key discussion points, decisions made, and important takeaways.

    Transcript:
    {transcript}
    """
    try:
        response = client.chat.completions.create(
            model="llama3-70b-8192",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        st.error(f"Error summarizing transcript: {e}")
        return "Summarization failed."

def extract_action_items(transcript):
    """Use Groq API to extract action items, responsible persons, and deadlines."""
    prompt = f"""
    You are an AI that extracts structured action items from meeting transcripts.
    Identify action items, responsible persons, and deadlines where available.
    Return the output as a JSON list with these keys: 'person', 'action', and 'deadline'.

    Ensure the deadline is extracted accurately. If no deadline is mentioned, leave it empty.

    Transcript:
    {transcript}
    """
    try:
        response = client.chat.completions.create(
            model="llama3-70b-8192",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3
        )
        action_items = json.loads(response.choices[0].message.content)

        # Parse dates using dateparser
        for item in action_items:
            if item["deadline"]:
                parsed_date = dateparser.parse(item["deadline"])
                item["deadline"] = parsed_date.strftime("%Y-%m-%d") if parsed_date else "Unclear deadline"
    except Exception as e:
        st.error(f"Error extracting action items: {e}")
        action_items = []
    return action_items

def generate_pdf(summary, action_items):
    """Generate a PDF report for download."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", style="B", size=16)
    pdf.cell(200, 10, "Meeting Summary Report", ln=True, align="C")
    pdf.set_font("Arial", size=12)
    pdf.ln(10)
    pdf.cell(200, 10, "Summary:", ln=True)
    pdf.multi_cell(0, 10, summary)
    pdf.ln(10)
    pdf.cell(200, 10, "Key Action Items:", ln=True)
    for item in action_items:
        pdf.multi_cell(0, 10, f"- {item['person']}: {item['action']} (Deadline: {item['deadline']})")
    pdf_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf.output(pdf_file.name)
    return pdf_file.name

# Streamlit UI
st.set_page_config(page_title="AI Meeting Summarizer", layout="wide", page_icon="📝")
st.markdown("<h1 style='text-align: center;'>📋 AI Meeting Summarizer</h1>", unsafe_allow_html=True)
st.markdown("<h4 style='text-align: center;'>Upload a meeting transcript</h4>", unsafe_allow_html=True)
st.markdown("---")

uploaded_file = st.file_uploader("Upload a transcript (TXT, PDF, DOCX)", type=["txt", "pdf", "docx"])
st.markdown("---")

if uploaded_file:
    transcript = extract_text_from_file(uploaded_file)
    st.subheader("📝 Transcript Preview")
    st.text_area("Transcript", transcript, height=200)

    if st.button("🚀 Generate Summary & Action Items"):
        summary = summarize_transcript(transcript)
        action_items = extract_action_items(transcript)
        
        st.subheader("📌 Summary")
        st.write(summary)
        
        st.subheader("✅ Key Action Items")
        if action_items:
            st.table(action_items)
        else:
            st.write("No clear action items detected.")
        
        pdf_path = generate_pdf(summary, action_items)
        with open(pdf_path, "rb") as pdf_file:
            st.download_button(label="📥 Download Summary (PDF)", data=pdf_file, file_name="Meeting_Summary.pdf", mime="application/pdf")
